"""
src/ui/components.py
---------------------
Reusable Streamlit UI components for Career Sensei.
"""

import streamlit as st
from src.rag.pipeline import AnalysisResult


# ── Helpers ───────────────────────────────────────────────────────────────────
def _bullets_html(text: str) -> str:
    """Convert newline-separated bullet text to an HTML <ul> list."""
    lines = [l.lstrip("-*• ").strip() for l in text.strip().split("\n") if l.strip()]
    if not lines:
        return "<p class='body'>—</p>"
    return "<ul>" + "".join(f"<li>{l}</li>" for l in lines) + "</ul>"


def _score_classes(score: int):
    """Return (circle_css_class, label_css_class, fit_label) for a given score."""
    if score >= 70:
        return "score-high", "score-high-txt", "Strong fit"
    if score >= 45:
        return "score-mid",  "score-mid-txt",  "Partial fit"
    return "score-low", "score-low-txt", "Stretch role"


# ── Components ────────────────────────────────────────────────────────────────
def section_header(title: str) -> None:
    st.markdown(
        f"<div style='font-size:11px;font-weight:700;letter-spacing:1.5px;"
        f"text-transform:uppercase;color:#9ca3af;margin:28px 0 14px;'>{title}</div>",
        unsafe_allow_html=True,
    )


def info_card(icon: str, label: str, color: str, body_html: str) -> None:
    """Render a bordered card with an icon label and arbitrary HTML body."""
    st.markdown(f"""
    <div style="background:#fff;border:1px solid #e8eaf0;border-radius:12px;padding:22px 24px;height:100%;box-sizing:border-box;">
      <div style="font-size:11px;font-weight:700;letter-spacing:1.2px;text-transform:uppercase;margin-bottom:10px;display:flex;align-items:center;gap:6px;color:{color};">{icon} {label}</div>
      {body_html}
    </div>""", unsafe_allow_html=True)


def score_card(score: int) -> None:
    """Render the circular match score badge card with inline gradient styles."""
    if score >= 70:
        grad = "linear-gradient(135deg,#10b981,#059669)"
        shadow = "0 4px 20px rgba(16,185,129,.25)"
        label_color = "#059669"
        fit_label = "Strong fit"
    elif score >= 45:
        grad = "linear-gradient(135deg,#f59e0b,#d97706)"
        shadow = "0 4px 20px rgba(245,158,11,.25)"
        label_color = "#d97706"
        fit_label = "Partial fit"
    else:
        grad = "linear-gradient(135deg,#ef4444,#dc2626)"
        shadow = "0 4px 20px rgba(239,68,68,.25)"
        label_color = "#dc2626"
        fit_label = "Stretch role"

    st.markdown(f"""
    <div style="background:#fff;border:1px solid #e8eaf0;border-radius:12px;text-align:center;padding:28px 16px;">
      <div style="text-align:center;">
        <div style="display:inline-flex;align-items:center;justify-content:center;width:110px;height:110px;border-radius:50%;font-size:36px;font-weight:700;color:#fff;margin-bottom:6px;background:{grad};box-shadow:{shadow};">{score}</div>
        <div style="font-size:12px;font-weight:600;letter-spacing:0.8px;text-transform:uppercase;color:{label_color};">{fit_label}</div>
      </div>
    </div>""", unsafe_allow_html=True)


def results_layout(analysis: AnalysisResult) -> None:
    """
    Render the full structured analysis output:
      Row 1 — Score  |  Job Summary
      Row 2 — Strengths  |  Gaps
      Row 3 — Experience Suggestions (full-width)
    """
    score = analysis.score or 0

    # Row 1: Score + Job Summary
    section_header("Overview")
    col_score, col_summary = st.columns([1, 3], gap="medium")
    with col_score:
        score_card(score)
    with col_summary:
        info_card("📋", "What This Role Is About", "#6366f1",
                  f"<p class='body'>{analysis.job_summary}</p>")

    # Candidate-friendly explanation of the score
    section_header("Match Score Explained")
    # Important: keep this HTML string unindented so Streamlit doesn't treat it as a code block.
    score_explainer_html = (
        "<p class='body'>"
        "The match score is an AI estimate of how closely your resume aligns to the job requirements."
        " It looks for evidence of fit like:"
        "</p>"
        "<ul>"
        "<li>Relevant skills, tools, and keywords mentioned in the job</li>"
        "<li>Experience level (similar responsibilities and outcomes)</li>"
        "<li>Depth/seniority signals (how much you’ve done vs. just learned)</li>"
        "<li>Any obvious missing areas compared to what the role asks for</li>"
        "</ul>"
        "<p class='body'>"
        "<b>How to read it:</b> <b>70+</b> strong fit, <b>45–69</b> partial fit, <b>&lt;45</b> stretch role (focus on the “Key Gaps” + “Action Plan”)."
        "</p>"
    )
    info_card(
        "🧠",
        "How we calculate the 0-100 score",
        "#6366f1",
        score_explainer_html,
    )

    # Row 2: Strengths + Gaps
    section_header("Fit Analysis")
    col_str, col_gap = st.columns(2, gap="medium")
    with col_str:
        info_card("✅", "Key Strengths", "#059669", _bullets_html(analysis.strengths))
    with col_gap:
        info_card("🔍", "Key Gaps", "#dc2626", _bullets_html(analysis.gaps))

    # Row 3: Experience Suggestions
    section_header("Action Plan")
    info_card("💡", "How to Get the Missing Experience", "#d97706",
              _bullets_html(analysis.suggestions))


def chatbot_section(chat_chain) -> None:
    """
    Render the follow-up chatbot section.
    Uses st.chat_input() so Enter key submits naturally.
    Reads/writes st.session_state.chat_history.
    """
    st.divider()
    st.html("<h3 style='color:#1a1d2e;font-size:18px;font-weight:600;margin-bottom:0;'>💬 Ask a Follow-up Question</h3>")
    st.html(
        "<p style='font-size:13px;color:#9ca3af;margin-top:4px;'>"
        "Ask anything about this role, your resume, or next steps.</p>"
    )

    # Render history oldest-first
    for role, msg in st.session_state.chat_history:
        cls = "bubble-user" if role == "user" else "bubble-ai"
        st.markdown(f'<div class="{cls}">{msg}</div>', unsafe_allow_html=True)

    # st.chat_input natively submits on Enter and returns the text
    chat_input = st.chat_input("e.g. What skill should I prioritize learning first?")
    if chat_input:
        try:
            with st.spinner("Thinking..."):
                from src.rag.pipeline import run_chat
                answer = run_chat(chat_chain, chat_input)
            st.session_state.chat_history.append(("user", chat_input))
            st.session_state.chat_history.append(("ai", answer))
            st.rerun()
        except Exception as e:
            err = str(e).lower()
            st.session_state.chat_history.append(("user", chat_input))
            if "429" in err or "quota" in err or "resourceexhausted" in err.replace(" ", ""):
                st.session_state.chat_history.append(("ai", "We've hit our usage limit. Please wait a minute and try again."))
            else:
                st.session_state.chat_history.append(("ai", "Something went wrong. Please try again in a moment."))
            st.rerun()


def learning_plan_tab(chat_chain, analysis: AnalysisResult) -> None:
    """
    Generate and display a simple, structured learning plan.

    The plan is generated on-demand (button click) and can be exported as a
    Word document.
    """

    # Session-state keys (set if missing so this function can be called safely)
    if "learning_plan_markdown" not in st.session_state:
        st.session_state.learning_plan_markdown = None
    if "learning_plan_weeks" not in st.session_state:
        st.session_state.learning_plan_weeks = None
    if "learning_plan_hours" not in st.session_state:
        st.session_state.learning_plan_hours = None
    if "learning_plan_docx_bytes" not in st.session_state:
        st.session_state.learning_plan_docx_bytes = None

    st.divider()
    section_header("Learning Plan")

    col_a, col_b = st.columns(2, gap="medium")
    with col_a:
        weeks = st.number_input(
            "Duration (weeks)",
            min_value=1,
            max_value=24,
            value=int(st.session_state.learning_plan_weeks or 4),
            step=1,
        )
    with col_b:
        hours_per_week = st.number_input(
            "Hours/week",
            min_value=1,
            max_value=40,
            value=int(st.session_state.learning_plan_hours or 5),
            step=1,
        )

    generate = st.button("Generate Learning Plan", use_container_width=True)

    def _score_label(score: int) -> str:
        if score >= 70:
            return "strong fit"
        if score >= 45:
            return "partial fit"
        return "stretch role"

    def _markdown_to_docx_bytes(markdown_text: str, doc_title: str) -> bytes:
        """
        Convert a restricted markdown subset into a Word document.

        Supported:
        - Headings: # / ## / ###
        - Bullets: lines starting with '- '
        - Paragraphs: everything else
        """
        from io import BytesIO
        from docx import Document
        from docx.shared import Pt

        doc = Document()

        # Make it look professional and consistent.
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

        doc.add_heading(doc_title, level=0)

        score = analysis.score if analysis.score is not None else 0
        doc.add_paragraph(f"Match score: {score}/100 ({_score_label(score)})")
        if analysis.job_summary:
            doc.add_paragraph(f"Role summary: {analysis.job_summary}")

        doc.add_paragraph("")  # spacer

        for raw_line in markdown_text.splitlines():
            line = raw_line.strip()
            if not line:
                continue

            # Headings
            if line.startswith("# "):
                doc.add_heading(line[2:].strip(), level=0)
                continue
            if line.startswith("## "):
                doc.add_heading(line[3:].strip(), level=1)
                continue
            if line.startswith("### "):
                doc.add_heading(line[4:].strip(), level=2)
                continue

            # Bullets
            if line.startswith("- "):
                doc.add_paragraph(f"• {line[2:].strip()}")
                continue

            # Paragraphs (strip simple bold markers)
            cleaned = line.replace("**", "")
            doc.add_paragraph(cleaned)

        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()

    if generate:
        # In deployed environments, the model call sometimes fails; show details
        # so users never get stuck with a generic message.
        if not getattr(analysis, "gaps", "").strip():
            st.error("Missing gap information. Please run analysis again.")
            return

        prompt = f"""You are a career coach. Create a simple, structured learning plan based on the candidate's resume-vs-job gap analysis.

Match score: {analysis.score}/100 ({_score_label(analysis.score or 0)}).
Duration requested: {weeks} week(s).
Time available: about {hours_per_week} hours per week.

Top gaps:
{analysis.gaps}

Experience suggestions (ideas to close gaps):
{analysis.suggestions}

Constraints:
- Keep it candidate-friendly and practical.
- Make it easy to follow (clear weekly blocks).
- Each week should include:
  1) Focus (1 sentence)
  2) Tasks (2-4 bullets)
  3) Deliverables (1-3 bullets that prove progress)
  4) Resume updates (1-2 copy/paste bullets)
- Output in Markdown with exactly this structure:

# Learning Plan for this Role
## Top 3 Gaps to Prioritize
- ...
## Learning Roadmap ({weeks} weeks)
### Week 1
**Focus:** ...
**Tasks:**
- ...
**Deliverables (done looks like):**
- ...
**Resume updates (copy/paste bullets):**
- ...
### Week 2
(repeat for each week up to Week {weeks})

Important: Do not include any extra commentary outside the Markdown structure.
"""

        try:
            from src.rag.pipeline import run_chat

            with st.spinner("Generating your learning plan..."):
                plan_md = run_chat(chat_chain, prompt)

            st.session_state.learning_plan_markdown = plan_md
            st.session_state.learning_plan_weeks = weeks
            st.session_state.learning_plan_hours = hours_per_week
            st.session_state.learning_plan_docx_bytes = None  # regenerate on download

        except Exception as e:
            st.error("Something went wrong while generating the learning plan.")
            with st.expander("Debug details (copy/paste this)"):
                st.write(f"{type(e).__name__}: {e}")

    if st.session_state.learning_plan_markdown:
        st.markdown(st.session_state.learning_plan_markdown)

        want_new = (
            st.session_state.learning_plan_weeks != weeks
            or st.session_state.learning_plan_hours != hours_per_week
        )
        if want_new:
            st.caption("Tip: click “Generate Learning Plan” again to match your new weeks/hours.")

        if st.session_state.learning_plan_docx_bytes is None:
            st.session_state.learning_plan_docx_bytes = _markdown_to_docx_bytes(
                st.session_state.learning_plan_markdown,
                "Career Sensei Learning Plan",
            )

        st.download_button(
            label="Download as Word (.docx)",
            data=st.session_state.learning_plan_docx_bytes,
            file_name="Career_Sensei_Learning_Plan.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )


def landing_page() -> None:
    st.html("""
    <div style="text-align:center; padding:90px 20px 0;">
      <div style="font-size:52px;margin-bottom:18px;">🥷</div>
      <h1 style="font-size:28px;font-weight:700;color:#1a1d2e;margin-bottom:10px;">Career Sensei</h1>
      <p style="color:#6b7280;font-size:15px;max-width:400px;margin:0 auto;line-height:1.8;">
        Upload your resume and paste a job URL to get an instant gap analysis,
        match score, and personalised action plan.
      </p>
    </div>""")


def waiting_card() -> None:
    st.markdown("""
    <div style="background:#fff;border:1px solid #e8eaf0;border-radius:12px;text-align:center;padding:48px;margin-top:20px;">
      <div style="font-size:32px;margin-bottom:12px;">📄</div>
      <p style="font-size:14px;color:#6b7280;">
        Resume uploaded. Paste a job URL in the sidebar and click <strong>Analyze</strong>.
      </p>
    </div>""", unsafe_allow_html=True)
